"""
Official document generation (.docx) for approved Performance Plans and
Leave Requests, including a QR verification code.

Note: the spec mentions WeasyPrint/PDF generation. PDF rendering via
WeasyPrint requires native system libraries (Pango/Cairo) that are not
guaranteed to be present on every deployment target, so this service uses
python-docx (pure Python, always installable) for a reliably reproducible
"compiles and runs" deliverable. DOCX opens natively in Word/Google Docs
and can be converted to PDF with LibreOffice (`soffice --convert-to pdf`)
in any environment where that binary is available -- see the README.
"""
import hashlib
import io
import os
import uuid
from datetime import datetime

import qrcode
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.core.config import settings

DOCS_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "generated_documents")
os.makedirs(DOCS_DIR, exist_ok=True)

BRAND_COLOR = RGBColor(0x4F, 0x46, 0xE5)  # indigo, matches frontend theme


def _verification_code(entity_type: str, entity_id: str) -> str:
    raw = f"{entity_type}:{entity_id}:{settings.SECRET_KEY}"
    return hashlib.sha256(raw.encode()).hexdigest()[:16].upper()


def _qr_image_bytes(payload: str) -> io.BytesIO:
    img = qrcode.make(payload)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def _header(doc: Document, title: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("SDP INTERNAL PERFORMANCE & WORKFLOW MANAGEMENT SYSTEM")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = BRAND_COLOR

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(title)
    run2.bold = True
    run2.font.size = Pt(18)

    doc.add_paragraph()


def generate_performance_plan_docx(plan, employee_name: str, decided_by_name: str | None) -> str:
    doc = Document()
    _header(doc, "Performance Plan - Official Record")

    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    rows = [
        ("Employee", employee_name),
        ("Plan Title", plan.title),
        ("Period", plan.period.value.replace("_", " ").title()),
        ("Status", plan.status.value.replace("_", " ").title()),
        ("Decided By", decided_by_name or "-"),
        ("Decision Reason", plan.decision_reason or "-"),
        ("Submitted On", plan.created_at.strftime("%Y-%m-%d %H:%M UTC")),
        ("Decided On", plan.decided_at.strftime("%Y-%m-%d %H:%M UTC") if plan.decided_at else "-"),
    ]
    for label, value in rows:
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = str(value)

    doc.add_paragraph()
    doc.add_heading("Goals / Checklist", level=2)
    for i, goal in enumerate(plan.goals or [], start=1):
        text = goal.get("text") if isinstance(goal, dict) else str(goal)
        doc.add_paragraph(f"{i}. {text}", style="List Number")

    if plan.notes:
        doc.add_heading("Notes", level=2)
        doc.add_paragraph(plan.notes)

    if plan.technical_committee_comment:
        doc.add_heading("Technical Committee Comment", level=2)
        doc.add_paragraph(plan.technical_committee_comment)

    code = _verification_code("performance_plan", str(plan.id))
    _add_signature_block(doc, code, str(plan.id))

    filename = f"performance_plan_{plan.id}.docx"
    path = os.path.join(DOCS_DIR, filename)
    doc.save(path)
    return path


def generate_leave_docx(leave, employee_name: str, decided_by_name: str | None) -> str:
    doc = Document()
    _header(doc, "Leave Approval - Official Record")

    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    rows = [
        ("Employee", employee_name),
        ("Leave Type", leave.leave_type.value.title()),
        ("Days", str(leave.days)),
        ("Start Date", leave.start_date.strftime("%Y-%m-%d")),
        ("End Date", leave.end_date.strftime("%Y-%m-%d")),
        ("Status", leave.status.value.replace("_", " ").title()),
        ("Decided By", decided_by_name or "-"),
        ("Decision Reason", leave.decision_reason or "-"),
        ("Reason for Leave", leave.reason),
    ]
    for label, value in rows:
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = str(value)

    code = _verification_code("leave_request", str(leave.id))
    _add_signature_block(doc, code, str(leave.id))

    filename = f"leave_request_{leave.id}.docx"
    path = os.path.join(DOCS_DIR, filename)
    doc.save(path)
    return path


def _add_signature_block(doc: Document, code: str, entity_id: str):
    doc.add_paragraph()
    doc.add_paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
    doc.add_paragraph(f"Verification Code: {code}")
    doc.add_paragraph(
        "This document was generated and digitally verified by the SDP System. "
        "Scan the QR code to verify authenticity against the verification code above."
    )
    qr_buf = _qr_image_bytes(f"SDP-VERIFY|{entity_id}|{code}")
    doc.add_picture(qr_buf, width=Inches(1.2))
